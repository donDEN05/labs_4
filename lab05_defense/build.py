"""Build five reports and a 14-slide defense from recorded experiments."""
import csv
import hashlib
import json
from pathlib import Path
import shutil
import sys
import textwrap
from zipfile import ZipFile,ZIP_DEFLATED
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx import Presentation
from pptx.util import Inches,Pt as PptPt
from pptx.dml.color import RGBColor as PptColor
from shared.data import ROOT,DATA
from shared.metrics import save
from shared.protocol import check_analysis

TOPIC='Разработка многоагентной системы персонализированного управления скидками в электронной коммерции'
FOLDERS=['lab01_llm','lab02_multimodal','lab03_marl','lab04_optimization','lab05_defense']
TITLES=['Большие языковые модели','Мультимодальные модели','Мультиагентное обучение с подкреплением','Методы оптимизации и стохастические процессы','Защита исследовательского проекта']
HYPOTHESES=[
 'H1: текстовая LLM превосходит частотный прогноз покупки при одинаковом временном тесте.',
 'H2: слияние модальностей улучшает Macro-F1 относительно сильнейшей одномодальной схемы.',
 'H3: координация агентов увеличивает прибыль относительно равномерного распределения бюджета без нарушений лимитов.',
 'H4: устойчивый оптимизатор сохраняет преимущество при переносе на другую схему признаков.'
]


def read(folder,file):return pd.read_csv(ROOT/folder/'results'/file)
def js(path):return json.loads(Path(path).read_text(encoding='utf-8'))
def num(value):
    if value is None or (isinstance(value,float) and np.isnan(value)):return '—'
    return f'{value:.4f}' if isinstance(value,(float,np.floating)) else str(value)


def paragraph(doc,text,style=None):return doc.add_paragraph(str(text),style)


def hyperlink(doc,url):
    p=doc.add_paragraph();link=OxmlElement('w:hyperlink')
    link.set(qn('r:id'),doc.part.relate_to(url,RT.HYPERLINK,is_external=True))
    run=OxmlElement('w:r');props=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'237C8B');props.append(color)
    size=OxmlElement('w:sz');size.set(qn('w:val'),'18');props.append(size);run.append(props)
    text=OxmlElement('w:t');text.text=url;run.append(text);link.append(run);p._p.append(link)


def table(doc,frame,columns,labels=None):
    t=doc.add_table(rows=1,cols=len(columns));t.style='Table Grid'
    for cell,label in zip(t.rows[0].cells,labels or columns):
        cell.text=label
        shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'16324F');cell._tc.get_or_add_tcPr().append(shade)
        for run in cell.paragraphs[0].runs:run.font.color.rgb=RGBColor(255,255,255);run.bold=True
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for i,row in frame[columns].iterrows():
        cells=t.add_row().cells
        for cell,value in zip(cells,row):
            cell.text=num(value)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(3)
                for r in p.runs:r.font.size=Pt(9)
    paragraph(doc,'')


def document(title):
    doc=Document();s=doc.sections[0]
    doc.core_properties.title=title;doc.core_properties.subject=TOPIC;doc.core_properties.author='Индивидуальный исследовательский проект'
    s.page_width=Cm(21);s.page_height=Cm(29.7)
    s.top_margin=s.bottom_margin=s.left_margin=s.right_margin=Cm(2)
    normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(12)
    normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.line_spacing=1.12
    for name in ['Title','Heading 1','Heading 2']:
        doc.styles[name].font.name='Times New Roman';doc.styles[name].font.color.rgb=RGBColor.from_string('16324F')
    doc.add_heading(title,0);paragraph(doc,'Прикладные задачи машинного обучения');paragraph(doc,TOPIC)
    paragraph(doc,'Индивидуальный исследовательский проект. Дата экспериментов: 22 октября 2026 г.')
    return doc


def common_data(doc,audit,manifest):
    doc.add_heading('Данные, очистка и временной протокол',1)
    paragraph(doc,f"Источник: Complete Journey, обработанная версия пакета completejourney. Транзакции: {audit['raw_rows']:,}; товарный каталог: 92 331 строка. Точные дубликаты транзакций: {audit['duplicates']}; исключённые некорректные строки: {audit['excluded']}. Пропусков в исходных полях транзакций не найдено. Служебные купонные записи и топливо исключены из выбора товарных категорий.")
    paragraph(doc,f"Основная выборка: {audit['selected_households']} покупателей, 8 категорий. Дополнительный перенос: {audit['transfer_households']} покупателей, отсутствующих в обучении. Единица наблюдения — покупатель, категория и неделя. Класс 1 означает хотя бы одну покупку в категории в целевую неделю. Панель содержит {audit['panel_rows']:,} объектов.")
    table(doc,pd.DataFrame(audit['by_split']),['cohort','split','label','n'],['Группа','Часть','Класс','Объекты'])
    paragraph(doc,'Признаки используют четыре недели до целевой. Части идут в порядке времени 60/20/20, пограничные окна удалены: последняя целевая неделя предыдущей части раньше первой недели истории следующей части. Временной индекс взят из набора, гражданский часовой пояс не выдумывается. Масштабирование и текстовый словарь обучены только на train. Идентификатор покупателя не подаётся в модель.')
    paragraph(doc,'Дисбаланс требует Macro-F1 и F1 класса покупки наряду с accuracy. Повторяющиеся числовые профили допустимы у независимых покупателей и сами по себе не доказывают утечку. Автоматическая проверка меняет будущие транзакции и требует неизменности прошлых признаков.')
    paragraph(doc,f"Версия источника: {manifest['revision']}. SHA-256 панели: {audit['sample_hash']}. Лицензия пакета указана как CC0; URL, контрольные суммы файлов и схемы находятся в shared/data/manifest.json и schemas.json. Источник: {manifest['source']}")
    profile=DATA/'profile_audit.json'
    if profile.exists():
        r=js(profile)
        paragraph(doc,f"Аудит текстовых профилей: уникальных {r['unique_texts']}, повторяющихся групп {r['repeated_text_groups']}, точных профилей в нескольких частях {r['exact_profiles_crossing_parts']}. Для фиксированных {r['near_test_sample']} test-объектов проверен ближайший train-профиль по TF-IDF cosine; сходство ≥0,97 у {r['near_matches']}, в том числе один и тот же покупатель у {r['near_same_household']}. Пустые/сходные истории могут повторяться у независимых покупателей; это ограничение различимости, а не доказательство утечки будущей метки. Ближайшие пары и описательные статистики сохранены в shared/data.")


def literature(doc,folder):
    refs=pd.read_csv(ROOT/folder/'literature.csv')
    doc.add_heading('Исследовательский обзор и разрыв',1)
    if folder!='lab05_defense':
        screening=js(ROOT/folder/'results/literature_screening.json')
        paragraph(doc,'Поисковый запрос: '+screening['queries'])
        paragraph(doc,f"PRISMA-подобный учёт: рассмотрено {screening['identified']}, дублей {screening['duplicates_removed']}, исключено {screening['excluded']}, включено {screening['included']}. Использован целевой поиск первичных работ с переходом по ссылкам. Это учёт реального ограниченного скрининга, а не заявление об исчерпывающем поиске во всех базах. Причины исключения и дата сохранены в literature_screening.json.")
    else:
        paragraph(doc,f'Итоговый синтез содержит {len(refs)} уникальных первичных источников из четырёх обзоров. По одному и тому же URL повторные записи объединены. Полная матрица содержит архитектуру, данные, протокол, метрики, код, ограничения и идею репликации.')
    for i,r in refs.iterrows():
        doc.add_heading(f"{i+1}. {r.authors} ({r.year})",2)
        paragraph(doc,f'{r.title}. {r.venue}. {r.architecture}. Исходные данные: {r.data}. Протокол: {r.protocol}. Метрики: {r.metrics}.')
        paragraph(doc,f'Ограничение для ВКР: {r.limitations}. Решение для эксперимента: {r.replication}.')
        hyperlink(doc,r.url)
    paragraph(doc,'Полная сравнительная матрица приложена в literature.csv. Рейтинги A/A* и квартиль Q1 не приписываются изданиям без отдельной проверки; статус препринтов и первичного лекционного источника RMSProp указан явно.')


def figure(doc,path,caption):
    if Path(path).exists():
        doc.add_picture(str(path),width=Cm(16))._inline.docPr.set('descr',caption)
        paragraph(doc,caption)


def plots():
    plt.rcParams.update({'font.family':'DejaVu Sans','font.size':10})
    specs=[('lab02_multimodal','scheme','macro_f1'),('lab03_marl','algorithm','profit_mean'),('lab04_optimization','optimizer','mean_f1')]
    for folder,name,value in specs:
        frame=read(folder,'metrics.csv').query("seed == 'ensemble' and cohort == 'main'") if folder=='lab02_multimodal' else read(folder,'summary.csv')
        frame=frame.sort_values(value)
        fig,ax=plt.subplots(figsize=(8,4));ax.barh(frame[name],frame[value],color='#237C8B');ax.set_xlabel('Macro-F1' if folder!='lab03_marl' else 'Прибыль за эпизод, условные денежные единицы')
        ax.grid(axis='x',alpha=.2);fig.tight_layout();fig.savefig(ROOT/folder/'results/comparison.png',dpi=160);plt.close(fig)
    for folder,label in [('lab02_multimodal','scheme'),('lab03_marl','algorithm'),('lab04_optimization','optimizer')]:
        curve=read(folder,'learning.csv');value='profit' if folder=='lab03_marl' else 'val_macro_f1'
        fig,ax=plt.subplots(figsize=(8,4))
        for name,g in curve.groupby(label):
            g=g.groupby('step')[value].mean();ax.plot(g.index,g.values,label=name)
        ax.set_xlabel('Шаги среды' if folder=='lab03_marl' else 'Шаги обучения');ax.set_ylabel('Прибыль эпизода' if folder=='lab03_marl' else 'Validation Macro-F1');ax.legend(ncol=3,fontsize=8);fig.tight_layout();fig.savefig(ROOT/folder/'results/learning.png',dpi=160);plt.close(fig)
    m=read('lab04_optimization','metrics.csv').query("cohort == 'main'")
    names=sorted(m.optimizer.unique());fig,ax=plt.subplots(figsize=(8,4))
    ax.boxplot([m[m.optimizer==name].macro_f1.to_numpy() for name in names]);ax.set_xticks(range(1,len(names)+1),names,rotation=25);ax.set_ylabel('Test Macro-F1, пять seed');fig.tight_layout();fig.savefig(ROOT/'lab04_optimization/results/distribution.png',dpi=160);plt.close(fig)


def report1(audit,manifest):
    folder='lab01_llm';doc=document('Лабораторная работа 1. '+TITLES[0])
    doc.add_heading('Задача и гипотеза',1);paragraph(doc,HYPOTHESES[0]);paragraph(doc,'Из истории покупок требуется оценить вероятность покупки в конкретной категории на следующей неделе. Модель не получает будущую скидку или истинную метку. Выходная вероятность служит признаком для прототипа, а не разрешением автоматически выдать скидку.')
    literature(doc,folder)
    paragraph(doc,'Исследовательский разрыв выбранного корпуса: общие языковые и рекомендательные benchmark не отвечают, окупается ли компактная LLM при прогнозе редкой покупки из короткой транзакционной истории. Требуется единый временной тест с калибровкой и стоимостью; отсутствие преимущества будет отрицательным результатом.')
    common_data(doc,audit,manifest)
    doc.add_heading('Модели и единый эксперимент',1)
    config=js(ROOT/folder/'config.json');mm=js(ROOT/folder/'results/model_manifest.json')
    table(doc,pd.DataFrame([dict(model=m,revision=mm.get(m,{}).get('revision','недоступна')[:12],license=mm.get(m,{}).get('license','не проверена')) for m in config['models']]),['model','revision','license'],['LLM','Ревизия','Лицензия'])
    paragraph(doc,'Условия: zero-shot и six-shot, две формулировки инструкции, идентичные объекты, температура 0. Основной декодер сравнивает logits цифр 0 и 1 после префикса JSON; вероятность нормируется только между двумя допустимыми цифрами. JSON формируется ограниченным декодером. Нулевая невалидность такого выхода является свойством ограничения, а не доказательством послушания модели. Отдельные свободные JSON-пробы не смешиваются с основными прогнозами.')
    paragraph(doc,'INT8-квантизация линейных слоёв одинакова для всех LLM, embeddings — float32. Полноточный пилот исключён из основной таблицы. CPU, 16 ГБ RAM; малый фиксированный тест выбран по хешу без доступа к меткам. Объёмы: 16 validation, 32 test, 16 новых покупателей. Шесть примеров из train сбалансированы по классу. Верхняя граница контекста 1024 токена. PEFT не выполнялась из-за ограниченного CPU-бюджета. Задержки измерены на данном компьютере без изолированного ресурсного стенда; фоновые загрузки могут влиять на время, поэтому ранжирование производительности ограничено.')
    metrics=read(folder,'metrics.csv');selection=js(ROOT/folder/'results/selection.json')
    doc.add_heading('Количественные результаты',1)
    winners=[]
    for model,g in metrics.query("split == 'validation' and cohort == 'main'").groupby('model'):
        eligible=g[g.invalid_rate<=.05];eligible=eligible[eligible.macro_f1>=eligible.macro_f1.max()-.01]
        best=eligible.sort_values(['brier','latency_s']).iloc[0]
        r=metrics[(metrics.model==model)&(metrics['mode']==best['mode'])&(metrics.instruction==best.instruction)&(metrics.split=='test')&(metrics.cohort=='main')].iloc[0]
        winners.append(r.to_dict())
    table(doc,pd.DataFrame(winners),['model','macro_f1','brier','ece','latency_s'],['LLM, режим выбран на validation','Macro-F1','Brier','ECE','с/объект'])
    paragraph(doc,f"Кандидатов с результатами: {metrics.model.nunique()} из 8. Выбор только по validation: {selection.get('model','нет допустимого кандидата')}, {selection.get('mode','')}, инструкция {selection.get('instruction','')}. Правило: invalid ≤5%, F1 в пределах 0,01 от максимума, затем Brier и latency. Все режимы и доверительные интервалы находятся в results/metrics.csv.")
    table(doc,read(folder,'baseline.csv'),['split','cohort','macro_f1','brier'],['Часть','Покупатели','Частотный Macro-F1','Частотный Brier'])
    paragraph(doc,'Частотный baseline: (число прошлых недель с покупкой + 1)/6, то есть сглаженная частота четырёх недель. Он использует только ту же историю и не настраивается на test.')
    probes=ROOT/folder/'results/free_json_probes.jsonl'
    if probes.exists():
        free=pd.read_json(probes,lines=True);free=free.groupby('model').agg(n=('valid','size'),valid=('valid','sum')).reset_index();free['invalid_rate']=1-free.valid/free.n
        table(doc,free,['model','n','invalid_rate'],['LLM','Свободные пробы','Доля невалидных'])
        paragraph(doc,'Четыре свободные пробы на модель используют один фиксированный объект в разных условиях. Это проверка формата, а не статистически надёжная оценка производственной частоты ошибок.')
    doc.add_heading('Ошибки, перенос и интеграция',1)
    paragraph(doc,'Отдельно сохранены ошибочные прогнозы и свободные ответы, включая markdown-ограждения, незавершённый JSON и лишние объяснения. Формат проверяется строго без подстановки правильной метки. В отчёте по основной задаче probability является условной вероятностью цифры, а не эмпирически откалиброванной вероятностью покупки. ECE и Brier измеряют именно это ограничение.')
    stability=ROOT/folder/'results/stability.csv'
    if stability.exists():table(doc,pd.read_csv(stability),['model','mode','instruction_flip_rate','mean_probability_change'],['LLM','Режим','Смена класса','Среднее |Δp|'])
    paragraph(doc,'Разбор по классам, длине, категории и языку сохранён в error_groups.csv. Все исходные текстовые описания англоязычные: русскоязычная проверка не имитируется переводом меток или выдуманными отзывами.')
    if selection.get('status')=='selected':
        rows=metrics[(metrics.model==selection['model'])&(metrics['mode']==selection['mode'])&(metrics.instruction==selection['instruction'])&(metrics.split=='test')]
        table(doc,rows,['cohort','n','macro_f1','f1_buy','ci_low','ci_high'],['Покупатели','n','Macro-F1','F1 покупки','CI 2,5%','CI 97,5%'])
    paragraph(doc,'Новый срез — покупатели, не участвовавшие в обучении компактных моделей; это перенос внутри исходного магазина, а не доказательство переноса на другую площадку. LLM оценивает спрос. В прототипе риск-контроллер отдельно проверяет бюджет и маржу. Общая обработка отсутствующей истории и низкой уверенности с переходом на baseline требует отдельной реализации; текущий опыт проверяет формат свободного ответа и ограничения бюджета/маржи.')
    doc.add_heading('Вывод и воспроизведение',1)
    paragraph(doc,'Малый тест имеет широкие интервалы и не позволяет объявить LLM универсально лучшей. Следующий эксперимент — увеличить временной и пользовательский тест, сравнить с откалиброванной табличной моделью и проверить русскоязычные реальные обращения. Полный ранжированный результат зависит от INT8 и формулировки инструкции; few-shot улучшение не предполагается заранее.')
    pair=ROOT/folder/'results/paired_baseline.json'
    if pair.exists():
        r=js(pair);paragraph(doc,f"H1, выбранная по validation LLM против частотного baseline: ΔMacro-F1={r['delta']:.4f}, 95% CI [{r['ci'][0]:.4f};{r['ci'][1]:.4f}], p={r['p']:.4f}. Локальное подтверждение: {'да' if r['confirmed_local'] else 'нет'}. Сравнение относится только к фиксированному малому тесту.")
    paragraph(doc,'Первый запуск 3B завершился нативной ошибкой до первого прогноза при квантизации; исправлено удержание старых слоёв в памяти, после чего запуск повторён. Сбой сохранён в failed_attempts.json. SHA-256 всех файлов весов сверены с LFS указателями закреплённых ревизий; журнал — weight_integrity.json.')
    paragraph(doc,'Команда: .venv/Scripts/python.exe lab01_llm/run.py. Матрица литературы, manifest весов, config.json, fewshot_ids.json, sample.jsonl, predictions.jsonl, metrics.csv, errors.csv и free_json_probes.jsonl приложены. Локальный архив проекта заменяет внешнюю публикацию; URL студенческого репозитория не создавался.')
    doc.save(ROOT/folder/'report.docx')


def report2(audit,manifest):
    folder='lab02_multimodal';doc=document('Лабораторная работа 2. '+TITLES[1]);doc.add_heading('Постановка',1);paragraph(doc,HYPOTHESES[1])
    literature(doc,folder);paragraph(doc,'Разрыв: преимущества крупных vision-language моделей не доказывают, что визуализация тех же четырёх точек истории полезнее числового представления. Проверяется именно вклад представления и способа слияния, без приписывания изображению новых сведений.')
    common_data(doc,audit,manifest)
    doc.add_heading('Мультимодальный объект и синхронизация',1)
    paragraph(doc,'Объект содержит прошлые количества, выручку и скидки за четыре недели; их детерминированное английское описание; изображение 3×16×16. Каждый канал — график одной величины. Текст кодируется TF-IDF до 128 признаков, числа масштабируются train-статистиками, изображение обрабатывает маленькая CNN. Каталог связывается по product_id как many-to-one. Это компактные схемы интеграции, а не репликация BLIP/CLIP.')
    paragraph(doc,'Восемь схем: text, numbers, image, early, late, text_image, numbers_image, all. Early соединяет сырые текстовые/числовые признаки; late усредняет logits двух обучаемых ветвей; остальные соединяют компактные embeddings. Бюджет каждой схемы: 300 шагов AdamW, batch 128, LR 0,001, пять seed. Маски отсутствующих модальностей применяются до кодировщиков и при объединении.')
    metrics=read(folder,'metrics.csv');main=metrics.query("seed == 'ensemble' and cohort == 'main'")
    table(doc,main,['scheme','macro_f1','f1_buy','brier','ci_low','ci_high'],['Схема','Macro-F1','F1 покупки','Brier','CI 2,5%','CI 97,5%'])
    resources=read(folder,'validation.csv').groupby('scheme')[['elapsed','latency_s','parameters']].mean().reset_index();resources['weight_kb']=resources.parameters*4/1024
    table(doc,resources,['scheme','elapsed','latency_s','weight_kb'],['Схема','Обучение, с','Вывод, с/объект','Веса, КиБ'])
    paragraph(doc,'Память таблицы — объём float32-весов, нижняя граница расхода процесса, не peak RSS. Измеренная batch-latency не равна задержке одиночного онлайн-запроса. Общие данные и TF-IDF-кэш в этот объём не включены. Общий класс хранит и вычисляет также неактивные ветви с замаскированными выходами; числа параметров и задержки не характеризуют специально оптимизированные одномодальные реализации.')
    figure(doc,ROOT/folder/'results/comparison.png','Рисунок 1. Тестовые результаты ансамблей пяти seed; выбор сделан по validation.')
    choice=js(ROOT/folder/'results/selection.json')['scheme'];permutation=read(folder,'permutation.csv');row=permutation[permutation.scheme==choice].iloc[0]
    paragraph(doc,f"По validation выбрана схема {choice}. На test разность с baseline {row.baseline}: {row.delta:.4f} Macro-F1. Парный permutation test: p={row.p:.4f}, после поправки Бонферрони p={row.p_bonferroni:.4f}. Следовательно, преимущество после учёта множественных сравнений не подтверждено на уровне 0,05. Интервал прироста в CSV построен кластерно по покупателям; он не скорректирован на множественный поиск.")
    doc.add_heading('Абляции, пропуски и перенос',1)
    table(doc,read(folder,'ablations.csv'),['missing','macro_f1','f1_buy','brier'],['Нет модальности','Macro-F1','F1 покупки','Brier'])
    table(doc,metrics.query("seed == 'ensemble' and cohort == 'transfer'"),['scheme','macro_f1','f1_buy','brier'],['Новые покупатели','Macro-F1','F1 покупки','Brier'])
    figure(doc,ROOT/folder/'results/learning.png','Рисунок 2. Средняя валидационная сходимость по seed.')
    paragraph(doc,'Ошибки редкого класса покупки анализируются в errors.csv; общий дисбаланс делает accuracy недостаточной. Изображение нормируется отдельно внутри каждого окна и теряет абсолютный масштаб: это известное ограничение. Маскирование при тесте не заменяет обучение на независимом наборе пропусков. Результаты не доказывают пользу фотографий товара или отзывов: их в исходных данных нет.')
    doc.add_heading('Вывод',1);paragraph(doc,'Усложнение до трёх модальностей не получает подтверждённого преимущества после коррекции множественных сравнений. Для прототипа сохраняется компактная выбранная схема; дальнейшая работа требует независимых timestamped-модальностей и более длинного отложенного периода.')
    paragraph(doc,'Запуск: .venv/Scripts/python.exe lab02_multimodal/run.py. В results доступны все пять seed, checkpoints, learning.csv, validation.csv, metrics.csv, predictions.csv, permutation.csv, ablations.csv и errors.csv. Паспорт обоих первичных наборов находится в shared/data/schemas.json; лицензия и версия — в manifest.json.')
    doc.save(ROOT/folder/'report.docx')


def report3(audit,manifest):
    folder='lab03_marl';doc=document('Лабораторная работа 3. '+TITLES[2]);doc.add_heading('Задача и гипотеза',1);paragraph(doc,HYPOTHESES[2])
    literature(doc,folder);paragraph(doc,'Разрыв выбранного обзора: совместимость credit assignment с бюджетом и минимальной маржой не следует из результатов игровых сред. Требуется сравнение с равномерной скидкой и одноагентным PPO при одинаковых взаимодействиях и явной модели спроса.')
    common_data(doc,audit,manifest)
    doc.add_heading('Спецификация среды',1)
    paragraph(doc,'Агентов восемь, по одному на категорию. Наблюдение агента: прогноз спроса компактной модели, оставшаяся доля бюджета, доля пройденного эпизода, текущий режим и one-hot категории. Состояние центрального critic — объединённые наблюдения. Действия: 0%, 5%, 10%, 15%, 20%. Исполнение локальной политики не использует critic. Одноагентный PPO получает всё состояние и выбирает вектор действий.')
    paragraph(doc,'Переход: p_i=clip(base_i + 1,2*d_i − 0,25*mean(d_others), 0,001, 0,999). Покупка — Bernoulli, не более одной единицы на категорию и неделю. Цена оценивается только по train. Базовая себестоимость 60% цены, бюджет эпизода 5 денежных единиц. Награда: прибыль/10 − 0,05×число предложенных нарушений. Эпизод заканчивается после доступных недель соответствующей части. Вся причинная реакция на скидку и себестоимость заданы предположениями.')
    paragraph(doc,'Контроллер сохраняет минимальную маржу 10%, максимальную скидку 20% и консервативный бюджет при покупке по одной единице всех категорий. Он учитывает повышенную себестоимость стресс-сценария. Предложенные и фактически применённые нарушения учитываются отдельно. Жадное уменьшение скидок при равенстве может зависеть от порядка категорий; это ограничение контроллера.')
    doc.add_heading('Алгоритмы и одинаковый бюджет',1)
    paragraph(doc,'IQL — локальный TD; VDN — сумма Q; QMIX — монотонный mixer с неотрицательными весами; QTRAN — joint TD и ограничения оптимальности. IAC использует локальный value; IPPO — локальный value и clipped objective; MAPPO — центральный value; COMA — action-conditioned critic и counterfactual baseline. Это компактные feed-forward варианты с разделяемыми локальными параметрами, не полные воспроизведения игровых benchmark.')
    paragraph(doc,'Каждый алгоритм и PPO обучены на 2000 взаимодействиях для каждого из пяти seed; LR 0,001, gamma 0,95. Q-методы используют replay и target networks, actor-critic — последовательные rollout. Бюджет взаимодействий равен, число градиентных обновлений различается по природе алгоритмов и записано в реализации. Равномерный baseline выбирает одинаковую доступную скидку и делит бюджет между оставшимися неделями. Подтверждения сходимости за этот короткий бюджет нет.')
    summary=read(folder,'summary.csv');table(doc,summary,['algorithm','profit_mean','profit_sd','discount_cost','applied_violations'],['Метод','Прибыль','SD по seed','Расход скидок','Нарушения'])
    figure(doc,ROOT/folder/'results/comparison.png','Рисунок 1. Средняя прибыль в симуляции; это не исторический causal backtest.')
    winner=js(ROOT/folder/'results/selection.json')['algorithm'];best=summary[summary.algorithm==winner].iloc[0];uniform=summary[summary.algorithm=='Uniform'].iloc[0]
    paragraph(doc,f"Выбор по validation: {winner}. На test средняя прибыль {best.profit_mean:.4f} против {uniform.profit_mean:.4f} у равномерного baseline, разность {best.profit_mean-uniform.profit_mean:.4f}. Небольшая разность и интервалы не обосновывают устойчивого экономического превосходства. У Uniform SD по обучающим seed равна нулю, поскольку baseline не обучается и использует одинаковые сценарии; это не означает отсутствие рыночной неопределённости.")
    pair=ROOT/folder/'results/paired_uniform.json'
    if pair.exists():
        r=js(pair);paragraph(doc,f"Парный bootstrap с независимым ресэмплингом осей seed и сценария: Δ={r['mean_difference']:.4f}, CI [{r['ci'][0]:.4f};{r['ci'][1]:.4f}]. Это интервал для проверенного симулятора; локальное подтверждение H3: {'да' if r['confirmed_local'] else 'нет'}.")
    episodes=read(folder,'episodes.csv')
    stress=episodes[(episodes.algorithm==winner)&((episodes.regime!='normal')|(episodes.ablation!='none'))].groupby(['regime','ablation']).agg(profit=('profit','mean'),expense=('discount_cost','mean'),violations=('applied_violations','sum')).reset_index()
    doc.add_heading('Стресс-тест и абляции',1);table(doc,stress,['regime','ablation','profit','violations'],['Режим','Абляция','Прибыль','Нарушения'])
    paragraph(doc,'Low demand уменьшает базовый спрос вдвое, low elasticity уменьшает эластичность до четверти, high cost повышает себестоимость на 15 п.п. Markov-switching использует P=[[0,85;0,15],[0,30;0,70]] и спрос 1 или 0,5 от базового. Абляции убирают персонализацию, общий бюджетный сигнал или заменяют кооперативную политику IQL. Последняя абляция является заменой метода, а не причинно изолированным удалением одного слоя.')
    figure(doc,ROOT/folder/'results/learning.png','Рисунок 2. Кривые прибыли обучения, усреднённые по seed.')
    doc.add_heading('Вывод и риски',1);paragraph(doc,'Жёсткие ограничения выполняются на проверенных траекториях, однако это не гарантия для другой модели спроса. Положительное отклонение прибыли не равно доказанному эффекту в магазине. Требуются более длинное обучение, больше независимых сценариев, анализ бюджетного контроллера и рандомизированный эксперимент до внедрения. Признаки дохода и демографии в персонализацию не включаются.')
    paragraph(doc,'Запуск: .venv/Scripts/python.exe lab03_marl/run.py. environment.json, training.csv, learning.csv, validation.csv, episodes.csv, violations.jsonl, summary.csv, risk_card.json и checkpoints приложены. Доходность/Sharpe финансовой методички заменены прибылью, расходом скидок, стабильностью и выполнением ограничений в согласованной теме ВКР.')
    doc.save(ROOT/folder/'report.docx')


def report4(audit,manifest):
    folder='lab04_optimization';doc=document('Лабораторная работа 4. '+TITLES[3]);doc.add_heading('Цель и гипотеза',1);paragraph(doc,HYPOTHESES[3]);literature(doc,folder)
    paragraph(doc,'Разрыв: теоретические скорости для выпуклой задачи и предположений о шуме не определяют победителя маленькой невыпуклой модели на дисбалансном retail-наборе. Проверяются отдельно скорость, валидация, обобщение и устойчивость.')
    common_data(doc,audit,manifest)
    doc.add_heading('Карта случайности и контролируемый эксперимент',1)
    paragraph(doc,'Источники случайности: начальные веса, порядок mini-batch, выбор seed, стохастичность исходных покупок и сценарии режимов спроса. Dropout в этой минимальной архитектуре отсутствует; его случайность не имитируется. При одном seed у всех оптимизаторов совпадают хеш начальных весов и хеш последовательности batch. Полный split также хешируется.')
    paragraph(doc,'Фиксированная модель — схема all лабораторной 2. Оптимизаторы: SGD, Momentum 0,9, Nesterov 0,9, AdaGrad, RMSprop, Adam, AdamW, Adadelta. У каждого одна сетка LR {0,0001;0,001;0,01}, tuning seed 11; после выбора LR — seed 11,22,33,44,55. 300 шагов, batch 128, constant scheduler, градиент обрезается до нормы 5. Остальные параметры штатного PyTorch; AdamW использует decoupled decay по умолчанию, то есть сравниваются стандартные конфигурации, а не изолированная формула обновления.')
    summary=read(folder,'summary.csv');table(doc,summary,['optimizer','mean_f1','sd_f1','ci_low','ci_high'],['Оптимизатор','Test Macro-F1','SD','CI 2,5%','CI 97,5%'])
    resources=read(folder,'runs.csv').groupby('optimizer').agg(elapsed=('elapsed','mean'),target=('time_to_target_s','mean'),parameters=('parameters','mean')).reset_index()
    table(doc,resources,['optimizer','elapsed','target','parameters'],['Метод','Время, с','До F1=0,65, с','Параметры'])
    paragraph(doc,'Прочерк времени до порога означает, что целевое качество не достигнуто. Память ниже — фактическое хранение тензоров после одного дополнительного обновления checkpoint seed 11; размеры весов проверены на всех пяти seed. Активации, runtime и peak RSS не включены. Этот профиль не заменяет качество основного эксперимента.')
    table(doc,read(folder,'memory.csv'),['optimizer','parameter_bytes','gradient_bytes','optimizer_state_bytes','total_tensor_bytes'],['Метод','Веса, байт','Градиенты','Состояния','Всего, байт'])
    figure(doc,ROOT/folder/'results/comparison.png','Рисунок 1. Тестовый Macro-F1 по пяти seed.')
    selected=js(ROOT/folder/'results/selection.json')['optimizer'];paragraph(doc,f'По validation и устойчивости выбран {selected}. Максимальный отдельный test-результат не используется для повторного выбора. Правило допускает не более одного нестабильного запуска; среди результатов в пределах 0,01 Macro-F1 приоритет меньшему SD, затем времени. Средние и интервалы по пяти seed описывают вариацию обучения на фиксированном test, а не всю неопределённость генеральной совокупности.')
    doc.add_heading('Чувствительность, шум и перенос',1)
    sensitivity=read(folder,'sensitivity.csv');table(doc,sensitivity.pivot(index='optimizer',columns='learning_rate',values='val_macro_f1').reset_index(),['optimizer',.0001,.001,.01],['Метод','LR 0,0001','LR 0,001','LR 0,01'])
    noise=read(folder,'gradient_noise.csv').groupby('optimizer')[['trace_covariance','noise_ratio']].mean().reset_index();table(doc,noise,['optimizer','trace_covariance','noise_ratio'],['Метод','След ковариации','Noise ratio'])
    transfer=read(folder,'transfer.csv').groupby(['optimizer','scheme']).agg(mean_f1=('macro_f1','mean'),sd=('macro_f1','std')).reset_index();table(doc,transfer,['optimizer','scheme','mean_f1','sd'],['Метод','Схема переноса','Macro-F1','SD'])
    pair=ROOT/folder/'results/paired_transfer.json'
    if pair.exists():
        r=js(pair);paragraph(doc,f"Перенос {r['optimizer']} против {r['control']}: средняя ΔMacro-F1={r['mean_difference']:.4f}, парный по seed CI [{r['ci'][0]:.4f};{r['ci'][1]:.4f}]. Подтверждение на данном фиксированном test: {'да' if r['confirmed_local'] else 'нет'}.")
    block=js(ROOT/folder/'results/block_bootstrap.json');paragraph(doc,f"Блочный bootstrap: блок {block['block_weeks']} недели, всего {block['blocks']} блока test; CI Macro-F1 {block['ci'][0]:.4f}–{block['ci'][1]:.4f}. Из-за малого числа временных блоков интервал ненадёжен для сильного вывода о будущем.")
    scenarios=read(folder,'scenarios.csv').groupby('scenario')[['macro_f1','brier']].mean().reset_index();table(doc,scenarios,['scenario','macro_f1','brier'],['Сценарий','Macro-F1','Brier'])
    paragraph(doc,'Markov-switching использует опубликованную в config матрицу, а не оценённую модель рынка. В сценарии меняются числовые ковариаты, истинные исторические метки остаются фиксированными: это стресс представления, не симуляция контрфактических покупок. Оценка градиентного шума использует восемь batch и не доказывает sub-Gaussian распределение. GBM/OU/GARCH не добавляются, поскольку согласованный дискретный Markov-switching уже выполняет требование о стохастическом сценарии.')
    figure(doc,ROOT/folder/'results/learning.png','Рисунок 2. Валидационная сходимость, средняя по seed.')
    figure(doc,ROOT/folder/'results/distribution.png','Рисунок 3. Распределение пяти фактических seed каждого оптимизатора; малая выборка ограничивает выводы.')
    doc.add_heading('Вывод',1);paragraph(doc,'AdamW даёт воспроизводимый выбранный компромисс в рамках проверенной LR-сетки; это не универсальное превосходство. Общая сетка может быть невыгодна Adadelta, а 300 шагов ограничивают вывод о сходимости. Перенос проверяется по приложенной таблице, без объявления малых пересекающихся различий доказанным выигрышем.')
    paragraph(doc,'Запуск: .venv/Scripts/python.exe lab04_optimization/run.py. Приложены tuning/sensitivity, пять seed, хеши инициализаций, кривые, оценка градиентного шума, transfer.csv, block_bootstrap.json, scenarios.csv и фактические статусы запусков.')
    doc.save(ROOT/folder/'report.docx')


def report5(audit,manifest):
    folder='lab05_defense';doc=document('Лабораторная работа 5. '+TITLES[4])
    doc.add_heading('Исследовательская постановка',1)
    paragraph(doc,'Проблема: единая скидка не учитывает различия спроса, а независимые акции категорий конкурируют за бюджет. Объект — процесс выдачи персонализированных предложений. Предмет — координация агентов и прогнозные признаки при ограничении бюджета и маржи. Цель — проверить полезность сложных методов относительно простых контролей на воспроизводимом ограниченном эксперименте.')
    for h in HYPOTHESES:paragraph(doc,h)
    paragraph(doc,'Методическая новизна проекта — единый протокол времени, покупателей, стоимости и ограничений для всех четырёх сравнений. Мировая научная новизна алгоритмов не заявляется. Практическая значимость — возможность отклонить ненужное усложнение до расходов на онлайн-эксперимент.')
    literature(doc,folder);common_data(doc,audit,manifest)
    doc.add_heading('Итоговая архитектура и интерфейсы',1)
    paragraph(doc,'Сбор транзакций и каталога; валидация product_id, скидок и времени; история покупателя; текстовый прогноз LLM; числовой/визуальный/текстовый модуль; агенты товарных категорий; ограничитель бюджета и маржи; журнал и объяснение человеку. Политики работают офлайн и не отправляют реальные предложения.')
    paragraph(doc,'Панель передаётся как id, household, category, week, history_start/end, text, numbers, label. LLM выдаёт вероятность покупки, мультимодальный модуль — вероятность того же события. Агенты получают прогнозы и остаток бюджета; выход — вектор дискретных скидок. Контроллер проверяет реализуемость и сохраняет исходный и применённый векторы.')
    paragraph(doc,'В демонстрации воспроизводится сохранённый реальный ответ выбранной LLM, выполняется числовой модуль и политика; для одной категории фиксированное late fusion усредняет две вероятности. Обучение MARL использовало выбранную компактную модель. Эффективность полного объединённого контура отдельно не измерялась; демонстрация проверяет интерфейсы, а не доказывает прирост прибыли от LLM.')
    doc.add_heading('Сводные результаты и неопределённость',1)
    for label,selection in [('LLM','lab01_llm'),('Мультимодальность','lab02_multimodal'),('MARL','lab03_marl'),('Оптимизация','lab04_optimization')]:
        s=js(ROOT/selection/'results/selection.json')
        paragraph(doc,label+': '+str(s.get('model',s.get('scheme',s.get('algorithm',s.get('optimizer','нет выбора'))))))
    llm=js(ROOT/'lab01_llm/results/selection.json');lm=read('lab01_llm','metrics.csv')
    lm=lm[(lm.model==llm['model'])&(lm['mode']==llm['mode'])&(lm.instruction==llm['instruction'])&(lm.split=='test')&(lm.cohort=='main')].iloc[0]
    mm=js(ROOT/'lab02_multimodal/results/selection.json')['scheme'];m=read('lab02_multimodal','metrics.csv').query("scheme == @mm and seed == 'ensemble' and cohort == 'main'").iloc[0]
    rl=js(ROOT/'lab03_marl/results/selection.json')['algorithm'];r=read('lab03_marl','summary.csv').query('algorithm == @rl').iloc[0]
    op=js(ROOT/'lab04_optimization/results/selection.json')['optimizer'];o=read('lab04_optimization','summary.csv').query('optimizer == @op').iloc[0]
    overview=pd.DataFrame([dict(stage='LLM',method=llm['model'],metric='Macro-F1',value=lm.macro_f1,ci=f'[{lm.ci_low:.4f};{lm.ci_high:.4f}]'),dict(stage='Модальности',method=mm,metric='Macro-F1',value=m.macro_f1,ci=f'[{m.ci_low:.4f};{m.ci_high:.4f}]'),dict(stage='MARL',method=rl,metric='Прибыль симуляции',value=r.profit_mean,ci=f'[{r.ci_low:.4f};{r.ci_high:.4f}]'),dict(stage='Оптимизация',method=op,metric='Macro-F1',value=o.mean_f1,ci=f'[{o.ci_low:.4f};{o.ci_high:.4f}]')])
    table(doc,overview,['stage','method','metric','value','ci'],['Этап','Метод','Показатель','Значение','95% CI'])
    table(doc,read('lab02_multimodal','metrics.csv').query("seed == 'ensemble' and cohort == 'main'"),['scheme','macro_f1','brier','ci_low','ci_high'],['Схема','Macro-F1','Brier','CI 2,5%','CI 97,5%'])
    table(doc,read('lab03_marl','summary.csv'),['algorithm','profit_mean','profit_sd','applied_violations'],['Метод','Прибыль симуляции','SD','Нарушения'])
    table(doc,read('lab04_optimization','summary.csv'),['optimizer','mean_f1','sd_f1'],['Оптимизатор','Macro-F1','SD'])
    paragraph(doc,'Таблицы исходных конфигураций и всех режимов LLM находятся в первой работе; сравнение схем использует ансамбли пяти seed, оптимизаторов — распределение seed, MARL — обучающие seed и общие случайные сценарии. Эти интервалы имеют разные единицы неопределённости и не сравниваются между собой как одна метрика.')
    hypothesis_rows=[]
    for label,path in [('H1','lab01_llm/results/paired_baseline.json'),('H3','lab03_marl/results/paired_uniform.json'),('H4','lab04_optimization/results/paired_transfer.json')]:
        if (ROOT/path).exists():
            r=js(ROOT/path);hypothesis_rows.append(dict(hypothesis=label,status='Локально подтверждена' if r['confirmed_local'] else 'Не подтверждена',scope='Фиксированный тест / заданный симулятор'))
    hypothesis_rows.append(dict(hypothesis='H2',status='Не подтверждена после поправки',scope='Восемь сравнений, Bonferroni'))
    table(doc,pd.DataFrame(hypothesis_rows),['hypothesis','status','scope'],['Гипотеза','Результат','Граница вывода'])
    doc.add_heading('Абляции и отрицательные результаты',1)
    paragraph(doc,'Мультимодальный прирост не сохраняет значимость после поправки на восемь сравнений. Некоторые сложные схемы хуже простого текста. Для MARL небольшая разность средней прибыли с равномерным baseline не доказывает экономическую полезность. У многих конфигураций LLM свободный JSON может отличаться от схемы; constrained decoder скрывает только формат, но не ошибочный прогноз. 300/2000 шагов не доказывают сходимости.')
    paragraph(doc,'Все фактически завершённые и прерванные пилоты различаются в журнале. Неподтверждённые гипотезы сохраняются в итогах. Ни одно число не взято из демонстрационных таблиц исходной методички. Историческая классификация и синтетическая реакция спроса имеют разные подписи.')
    doc.add_heading('Ответственность и границы применения',1)
    paragraph(doc,'Покупательские идентификаторы псевдонимизированы источником; они применяются только для группировки и не становятся признаком. Доход, возраст и другие демографические поля исключены. На реальной площадке необходимо ограничить доступ к транзакциям, срок хранения и цель обработки. Нельзя считать отсутствие чувствительных полей гарантией отсутствия дискриминации: категории покупок также могут быть прокси.')
    paragraph(doc,'Перед внедрением нужны согласованная политика допустимых различий предложений, проверка отказов и ручной контроль. LLM не распоряжается бюджетом и не изменяет цены. Сейчас реализованы проверки формата ответа, допустимости действий, бюджета и маржи; общий fallback для неизвестного товара и отсутствующей истории остаётся следующим этапом. Правовые выводы без отдельной проверки юрисдикции не делаются.')
    paragraph(doc,'Основные угрозы: перенос продуктовой розницы на e-commerce; один магазин и короткое окно; повторение пользователей между временными частями; малый LLM-test; общая первичная информация трёх модальностей; допущения о стоимости и эластичности; отсутствие рандомизации скидок; общие сценарии baseline; узкая LR-сетка и короткое обучение.')
    doc.add_heading('Вывод и следующий эксперимент',1)
    paragraph(doc,'Проект реализует минимальную исследовательскую систему и воспроизводимый способ сравнения. Результаты поддерживают сохранение простых baseline и жёсткого контроллера, но не дают оснований обещать рост реальных продаж. Следующий шаг — получить timestamped онлайн-историю показов предложений и рандомизированную вариацию скидки, провести более длинный out-of-time тест и сравнить инкрементальную прибыль с заранее заданными ограничениями.')
    paragraph(doc,'Материалы защиты: 14 слайдов defense.pptx, demo_scenario.md, demo.py, results/demo.json, experiment_passport.json, contributions.csv, literature.csv, artifact_manifest.json и artifacts.zip. Архив содержит код, отчёты, таблицы, логи и инструкции; веса и окружение устанавливаются отдельно. Ссылка на внешний репозиторий не выдумывается, сдаётся локальный архив.')
    doc.save(ROOT/folder/'report.docx')


def presentation():
    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
    blue=PptColor(22,50,79);teal=PptColor(35,124,139);gray=PptColor(65,72,80)
    selected={f:js(ROOT/f/'results/selection.json') for f in FOLDERS[:4]}
    h1=js(ROOT/'lab01_llm/results/paired_baseline.json');h3=js(ROOT/'lab03_marl/results/paired_uniform.json');h4=js(ROOT/'lab04_optimization/results/paired_transfer.json')
    verdict=lambda result:'подтверждена локально' if result['confirmed_local'] else 'не подтверждена'
    choice=selected['lab01_llm'];lm=read('lab01_llm','metrics.csv')
    score=lm[(lm.model==choice['model'])&(lm['mode']==choice['mode'])&(lm.instruction==choice['instruction'])&(lm.split=='test')&(lm.cohort=='main')].iloc[0].macro_f1
    base=read('lab01_llm','baseline.csv').query("split == 'test' and cohort == 'main'").iloc[0].macro_f1
    data=js(DATA/'audit.json')
    pages=[
      ('Персонализированные скидки', ['Многоагентная система с контролем бюджета','Исследовательский проект по теме ВКР','Четыре сравнения, единые данные и ограничения'],None),
      ('Проблема и постановка',['Категории конкурируют за общий бюджет скидок','Прогноз: покупка в категории в следующую неделю','Решение: дискретная скидка с ограничением маржи'],None),
      ('Четыре проверяемые гипотезы',['H1. LLM полезнее частоты прошлых покупок','H2. Fusion улучшает одномодальный baseline','H3. Кооперация увеличивает прибыль симуляции','H4. Устойчивый оптимизатор переносится'],None),
      ('Литература обосновывает контроль',['По 12 источников в четырёх обзорах','36 уникальных источников в итоговой матрице','Языковые и игровые benchmark не доказывают эффект скидки'],None),
      ('Реальные данные и модельные допущения',[f"{data['raw_rows']:,} транзакций в исходной версии",'200 покупателей и 8 категорий в основной панели','20 новых покупателей для переноса','Спрос при новой скидке и себестоимость заданы в симуляции'],None),
      ('Временной тест исключает будущее',['Четыре недели истории до целевой покупки','Train, validation и test идут по времени','Пограничные окна удалены, словарь и масштабирование только train','Победитель фиксируется до test'],None),
      ('LLM: восемь моделей, один формат',[str(selected['lab01_llm'].get('model','Нет выбранной модели')),'Zero-shot и six-shot, две инструкции, INT8',f'Test Macro-F1: {score:.3f}; частотный контроль: {base:.3f}','Constrained JSON и отдельные свободные пробы'],None),
      ('Мультимодальность: эффект ограничен',['Восемь схем, обязательные text/numbers/image baseline','Выбрано: '+selected['lab02_multimodal']['scheme'],'Поправка на множественные сравнения снимает значимость прироста'],'lab02_multimodal/results/comparison.png'),
      ('MARL: ограничения важнее среднего',['Восемь алгоритмов, Uniform и одноагентный PPO','Выбрано: '+selected['lab03_marl']['algorithm'],'Прибыль относится к симуляции; преимущество над Uniform мало'],'lab03_marl/results/comparison.png'),
      ('Оптимизатор выбран по валидации',['Восемь методов, пять seed, одинаковые начальные веса','Выбрано: '+selected['lab04_optimization']['optimizer'],'Test максимум отдельного метода не меняет выбор'],'lab04_optimization/results/comparison.png'),
      ('Компоненты соединяет контроллер',['История и каталог; LLM и компактный прогноз','Агенты категорий предлагают скидки','Контроллер проверяет бюджет и минимальную маржу','Офлайн-демо воспроизводит сохранённый ответ LLM'],None),
      ('Отрицательные результаты сохранены',['Изображение производно от того же ряда','Проверены пропуски модальностей и отсутствие кооперации','Короткий бюджет не доказывает сходимости','Объединённый контур не проходил отдельный causal тест'],None),
      ('Риски и роль человека',['Один магазин, малый тест LLM, допущения спроса','Демография исключена, прокси-смещения всё ещё возможны','Реальные предложения и цены не изменяются','Перед внедрением нужен рандомизированный эксперимент'],None),
      ('Вывод: сложность требует доказательства',[f"H1: {verdict(h1)}; H2: нет после поправки",f"H3: {verdict(h3)}; H4: {verdict(h4)}",'H4 относится к пяти seed и одному фиксированному test','Следующий шаг: реальные показы предложений и A/B-протокол'],None)
    ]
    durations=[45,70,75,60,85,90,100,100,100,90,85,85,75,60]
    scenarios=[]
    for i,(title,bullets,picture) in enumerate(pages):
        slide=prs.slides.add_slide(prs.slide_layouts[6])
        box=slide.shapes.add_textbox(Inches(.65),Inches(.45),Inches(12),Inches(1.0))
        p=box.text_frame.paragraphs[0];p.text=title;p.font.name='Arial';p.font.size=PptPt(30);p.font.bold=True;p.font.color.rgb=blue
        width=5.0 if picture else 11.7
        body=slide.shapes.add_textbox(Inches(.75),Inches(1.8),Inches(width),Inches(4.8));tf=body.text_frame;tf.word_wrap=True
        for j,text in enumerate(bullets):
            p=tf.paragraphs[0] if j==0 else tf.add_paragraph();p.text=text;p.font.name='Arial';p.font.size=PptPt(22 if picture else 25);p.font.color.rgb=gray;p.space_after=PptPt(24)
        if picture:
            pic=slide.shapes.add_picture(str(ROOT/picture),Inches(6),Inches(2.05),width=Inches(6.8))
            pic._element.nvPicPr.cNvPr.set('descr',title+'. '+ ' '.join(bullets))
        f=slide.shapes.add_textbox(Inches(11.9),Inches(6.9),Inches(.7),Inches(.3));p=f.text_frame.paragraphs[0];p.text=f'{i+1}/14';p.font.size=PptPt(11);p.font.color.rgb=teal
        note=f'Время {durations[i]} с. '+ ' '.join(bullets)+'. '+[
            'Начать с цели исследования, а не с обещаний роста выручки. Обозначить индивидуальный формат и воспроизводимые результаты.',
            'Пояснить, почему одинаковая скидка может тратить бюджет на и без того вероятную покупку. Целевой эффект скидки отделён от вероятности покупки.',
            'Подчеркнуть, что каждая гипотеза может быть отвергнута и что отрицательный результат остаётся в защите.',
            'Показать матрицу с ограничениями и идеями репликации. Не называть препринты журнальными публикациями.',
            'Открыть manifest и паспорт. Показать контрольные суммы, исключённые строки и происхождение двух связанных наборов.',
            'Нарисовать временные границы. Показать тест изменения будущих транзакций и отсутствие будущих полей в истории.',
            'Пояснить нормирование logits цифр и структурную валидность JSON. Показать один свободный ответ и фактическую latency.',
            'Прочитать поправленное p-value. Пояснить, почему одинаковый первичный ряд в трёх представлениях не создаёт новые данные.',
            'Расписать награду и hard projection. Сравнить прибыль с Uniform, указать размер различия и ограничения доверительного интервала.',
            'Показать совпадающие хеши инициализации и порядок mini-batch. Разделить tuning seed и пять финальных seed.',
            'Запустить demo.py. Показать исходный и применённый вектор скидок, бюджет и минимальную маржу. Новых предложений покупателям нет.',
            'Открыть таблицы абляций и стресс-тестов. Объяснить, почему короткий training не доказывает превосходства алгоритмического семейства.',
            'Обсудить смещения, перенос между магазинами и отсутствие причинной идентификации. Пояснить требования к следующему эксперименту.',
            'Вернуться к четырём гипотезам и ограниченному выводу. Передать слово для вопросов; держать подготовленными исходные таблицы и архив.'
        ][i]
        slide.notes_slide.notes_text_frame.text=note
        scenarios.append(f'## Слайд {i+1}. {title} — {durations[i]} с\n\n{note}\n')
    prs.save(ROOT/'lab05_defense/defense.pptx')
    questions='''\n## До 10 минут вопросов\n\n1. Почему нельзя считать покупку доказательством эффекта скидки? Нет рандомизированного контрфактического исхода.\n2. Почему график не является независимыми данными? Он построен из того же прошлого ряда.\n3. Почему выбранный метод отличается от максимума test? Выбор зафиксирован по validation до test.\n4. Что гарантирует контроллер? Бюджет и маржу в оговорённой one-unit среде, а не в произвольном реальном магазине.\n5. Что означает нулевой invalid rate? Свойство constrained decoder; свободные пробы показываются отдельно.\n6. Почему только компактные модели? CPU и 16 ГБ RAM; пределы выборки и квантизация явно записаны.\n7. Где научная новизна? Авторский протокол и анализ ограничений; новые мировые алгоритмы не заявляются.\n8. Доказан ли выигрыш всей системы? Нет, компоненты оценены отдельно; demo проверяет взаимодействие.\n9. Что нужно для внедрения? Реальные показы предложений, рандомизация, контроль смещений и согласованные ограничения.\n10. Как воспроизвести числа? Запустить скрипты по README, использовать зафиксированные данные, ревизии и seed.\n'''
    (ROOT/'lab05_defense/demo_scenario.md').write_text(f'# Сценарий защиты\n\nОбщее время доклада: {sum(durations)//60} мин {sum(durations)%60} с. Вопросы: до 10 мин.\n\n'+ '\n'.join(scenarios)+questions,encoding='utf-8')


def supporting():
    audit=js(DATA/'audit.json');manifest=js(DATA/'manifest.json')
    datacard=dict(source=manifest,quality=audit,target='purchase in category in next week',unit='household/category/week',history='four previous weeks',limitations=['single grocery retailer','observational discounts','no randomized offer exposure','text/chart derived from same history'])
    for i,folder in enumerate(FOLDERS):
        save(ROOT/folder/'data_card.json',datacard)
        (ROOT/folder/'README.md').write_text(f'# Лабораторная {i+1}. {TITLES[i]}\n\n{TOPIC}\n\nОтчёт: report.docx. Литература: literature.csv. Паспорт: data_card.json. Конфигурация и фактические журналы: config.json и results/.\n\nЗапуск из корня проекта:\n\n```powershell\n.venv/Scripts/python.exe {folder}/'+('build.py' if i==4 else 'run.py')+f'\n.venv/Scripts/python.exe {folder}/test_lab.py\n```\n\nДанные и зависимости общие: ../shared и ../requirements.txt. Точная последовательность подготовки находится в корневом README.md. Полноценное повторение LLM требует загрузки весов; smoke-проверки не заменяют эксперимент.\n',encoding='utf-8')
    passport=dict(topic=TOPIC,date='2026-10-22',seeds=[11,22,33,44,55],data_revision=manifest['revision'],data_hash=audit['sample_hash'],configs={f:js(ROOT/f/'config.json') for f in FOLDERS[:4]},selection={f:js(ROOT/f/'results/selection.json') for f in FOLDERS[:4]},code_validation='eight runnable checks plus offline demo; final results recorded in verification.json',render_validation='Native Word/PowerPoint renderer unavailable on this Windows host; structural checks and source plots only')
    save(ROOT/'lab05_defense/experiment_passport.json',passport)
    pd.DataFrame([dict(participant='Индивидуальный проект; автоматизированное выполнение',role=role,hypothesis=h,artifact=artifact) for role,h,artifact in [
        ('Постановка и проверка ограничений','H1–H4','docs/superpowers/specs/2026-10-22-personalized-discounts-design.md'),
        ('Аудит данных и split','Отсутствие будущего в признаках','shared/data/audit.json'),
        ('LLM-эксперимент','H1','lab01_llm/results/predictions.jsonl'),
        ('Схемы модальностей','H2','lab02_multimodal/results/permutation.csv'),
        ('MARL и бюджетный контроллер','H3','lab03_marl/results/episodes.csv'),
        ('Контролируемая оптимизация','H4','lab04_optimization/results/runs.csv'),
        ('Отчёты и офлайн-демонстрация','Интеграция компонентов','lab05_defense/demo.py')]]).to_csv(ROOT/'lab05_defense/contributions.csv',index=False,encoding='utf-8-sig')
    save(ROOT/'lab05_defense/config.json',dict(slides=14,presentation_seconds=1120,questions_minutes=10,offline_only=True))


def archive():
    selected=[]
    for folder in FOLDERS+['shared','docs']:
        for path in (ROOT/folder).rglob('*'):
            rel=path.relative_to(ROOT)
            if not path.is_file() or any(p in ('__pycache__','models') for p in rel.parts):continue
            if path.suffix in ('.log','.partial','.bin','.headers','.rejected_sha256','.rds','.rda') or path.name in ('artifacts.zip','artifact_manifest.json','STOP'):continue
            if rel.parts[0]=='shared' and rel.parts[1:] and rel.parts[1] not in ('data',) and path.suffix not in ('.py','.json'):continue
            selected.append(path)
    selected.extend([ROOT/'README.md',ROOT/'requirements.txt',*ROOT.glob('*.docx')])
    entries=[dict(path=str(p.relative_to(ROOT)).replace('\\','/'),bytes=p.stat().st_size,sha256=hashlib.sha256(p.read_bytes()).hexdigest()) for p in selected]
    save(ROOT/'lab05_defense/artifact_manifest.json',dict(files=entries,exclusions=['model weights','virtual environment','raw R containers; reload using manifest','failed download files','temporary logs']))
    with ZipFile(ROOT/'lab05_defense/artifacts.zip','w',ZIP_DEFLATED) as z:
        for p in selected:z.write(p,p.relative_to(ROOT))
        z.write(ROOT/'lab05_defense/artifact_manifest.json','lab05_defense/artifact_manifest.json')


if __name__=='__main__':
    if '--archive-only' in sys.argv:
        archive();sys.exit(0)
    check_analysis(ROOT)
    audit=js(DATA/'audit.json');manifest=js(DATA/'manifest.json')
    plots()
    for writer in [report1,report2,report3,report4,report5]:writer(audit,manifest);print(writer.__name__,'written',flush=True)
    presentation();supporting();archive()
    print('Reports, presentation and archive written; run test_lab.py for evidence',flush=True)
