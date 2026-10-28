"""Checks final deliverables against raw logs; run after build.py."""
import json
from pathlib import Path
import sys
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import pandas as pd
from pptx import Presentation
from docx import Document
from shared.metrics import classification
from shared.protocol import check_analysis,digest
from lab01_llm.run import validate_records
from lab01_llm.run import CONFIG
import hashlib
from zipfile import ZipFile

root=Path(__file__).resolve().parents[1]
deck=Presentation(root/'lab05_defense/defense.pptx')
assert len(deck.slides)==14
for slide in deck.slides:
    assert len(slide.notes_slide.notes_text_frame.text)>100
    for shape in slide.shapes:
        assert 0<=shape.left and 0<=shape.top and shape.left+shape.width<=deck.slide_width and shape.top+shape.height<=deck.slide_height
        if shape.shape_type==13:assert shape._element.nvPicPr.cNvPr.get('descr')
for folder in ['lab01_llm','lab02_multimodal','lab03_marl','lab04_optimization','lab05_defense']:
    doc=Document(root/folder/'report.docx')
    assert len(doc.paragraphs)>15,folder
    assert all(shape._inline.docPr.get('descr') for shape in doc.inline_shapes),folder
    refs=pd.read_csv(root/folder/'literature.csv')
    assert refs.url.nunique()==len(refs)
    assert (25<=len(refs)<=40) if folder=='lab05_defense' else (10<=len(refs)<=15)
metrics=pd.read_csv(root/'lab02_multimodal/results/metrics.csv')
assert metrics.scheme.nunique()==8
pred=pd.read_csv(root/'lab02_multimodal/results/predictions.csv')
for scheme,g in pred[pred.cohort=='main'].groupby('scheme'):
    r=classification(g.label.to_numpy(),pd.DataFrame({'no':1-g.p,'buy':g.p}).to_numpy())
    saved=metrics[(metrics.scheme==scheme)&(metrics.seed=='ensemble')&(metrics.cohort=='main')].iloc[0]
    assert abs(r['macro_f1']-saved.macro_f1)<1e-10
marl=pd.read_csv(root/'lab03_marl/results/training.csv')
cfg=json.loads((root/'lab03_marl/config.json').read_text(encoding='utf-8'))
assert len(marl)==45 and set(zip(marl.algorithm,marl.seed))=={(a,s) for a in cfg['algorithms']+['PPO'] for s in cfg['seeds']} and (marl.status=='complete').all()
opt=pd.read_csv(root/'lab04_optimization/results/runs.csv')
cfg=json.loads((root/'lab04_optimization/config.json').read_text(encoding='utf-8'))
assert len(opt)==40 and set(zip(opt.optimizer,opt.seed))=={(a,s) for a in cfg['optimizers'] for s in cfg['seeds']} and (opt.status=='complete').all()
fusion=pd.read_csv(root/'lab02_multimodal/results/validation.csv')
cfg=json.loads((root/'lab02_multimodal/config.json').read_text(encoding='utf-8'))
assert len(fusion)==40 and set(zip(fusion.scheme,fusion.seed))=={(a,s) for a in cfg['schemes'] for s in cfg['seeds']} and (fusion.status=='complete').all()
llm=pd.read_json(root/'lab01_llm/results/predictions.jsonl',lines=True)
validate_records(llm.to_dict('records'),require_all=True)
probes=[json.loads(s) for s in (root/'lab01_llm/results/free_json_probes.jsonl').read_text(encoding='utf-8').splitlines()]
keys=[(r['model'],r['mode'],r['instruction']) for r in probes]
assert len(keys)==len(set(keys))==32
assert set(keys)=={(m,mode,i) for m in CONFIG['models'] for mode in ('zero-shot','six-shot') for i in (0,1)}
check_analysis(root)
manifest=json.loads((root/'lab05_defense/artifact_manifest.json').read_text(encoding='utf-8'))
with ZipFile(root/'lab05_defense/artifacts.zip') as archive:
    assert json.loads(archive.read('lab05_defense/artifact_manifest.json'))==manifest
    for entry in manifest['files']:
        assert digest(root/entry['path'])==entry['sha256'],entry['path']
        assert hashlib.sha256(archive.read(entry['path'])).hexdigest()==entry['sha256'],entry['path']
print('PASS lab05: five reports, 14 slides, source counts, raw-metric consistency, all model/seed counts, manifest')
